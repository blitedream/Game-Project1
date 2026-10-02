from __future__ import annotations

import argparse
import hashlib
import os
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REFERENCE_SHA256 = "5737D118FDB80934D826E7AD5CAF9D412C8E38A47774A4ED8F386FE9477C35A1"
FONT_LATIN = "DengXian"
FONT_EAST_ASIA = "等线"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str | None = None) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_EAST_ASIA)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value and keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    elif not value and keep is not None:
        p_pr.remove(keep)


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_title(document: Document, text: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, 19, True)
    set_keep_with_next(paragraph)

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    run = meta.add_run(subtitle)
    set_run_font(run, 10, False, "555555")


def add_heading(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(11)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(f"{number}  {text}")
    set_run_font(run, 12.5, True)
    set_keep_with_next(paragraph)


def add_subheading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5, True)
    set_keep_with_next(paragraph)


def add_body(document: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 10.5, True)
        remainder = paragraph.add_run(text[len(bold_lead):])
        set_run_font(remainder)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28 + 0.30 * level)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.22
    paragraph.paragraph_format.keep_together = True
    marker = "•" if level == 0 else "◦"
    run = paragraph.add_run(f"{marker}  {text}")
    set_run_font(run)


def add_numbered_item(document: Document, number: int, title: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.24)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.24
    paragraph.paragraph_format.keep_together = True
    lead = paragraph.add_run(f"{number}. {title}：")
    set_run_font(lead, 10.5, True)
    body = paragraph.add_run(text)
    set_run_font(body)


def build_document(source: Path, output: Path) -> None:
    actual_hash = sha256(source)
    if actual_hash != REFERENCE_SHA256:
        raise RuntimeError(
            f"Reference changed. Expected {REFERENCE_SHA256}, got {actual_hash}."
        )

    if source == output:
        raise RuntimeError("Output must be different from the retained reference.")

    output.parent.mkdir(parents=True, exist_ok=True)
    working = output.with_name(f"{output.stem}.working.docx")
    package_tmp = output.with_name(f"{output.stem}.package.tmp.docx")
    shutil.copy2(source, working)
    document = Document(working)
    clear_body(document)

    section = document.sections[0]
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    add_title(
        document,
        "详细游戏设计文档",
        "工作标题：《深潜：洞穴救援》（暂定）  |  项目代号：GP1  |  文档版本：v0.9  |  制作人员：待填写",
    )

    add_heading(document, 1, "工作标题")
    add_body(document, "《深潜：洞穴救援》（Deep Dive: Cave Recovery，暂定）。标题同时指向洞穴潜水的核心动作、逐关加深的风险，以及第三关从训练转向真实打捞救援的叙事升级。")

    add_heading(document, 2, "当前状态")
    add_body(document, "项目处于可玩垂直切片阶段。三关、三维任务地图、关卡解锁和结算流程已建立；第一关流程相对完整，第二、三关的任务逻辑已可运行，环境模型、水域连通、碰撞与出生点仍在集中修整。")
    add_subheading(document, "已完成或已评审")
    add_bullet(document, "第一人称陆地与水下移动、鼠标观察、冲刺、上下潜、浮力、重力、水阻和八方向随机水流。")
    add_bullet(document, "三关 Enter 开始、教练对话、岸边领取装备、任务清单、计时结算、最佳时间和逐关解锁。")
    add_bullet(document, "第一关 60 m 训练池与 20 m / 40 m / 60 m 深度目标；完成后标记变灰并停止旋转。")
    add_bullet(document, "三维任务地图、节点状态、聚焦/缩放、M 键开图、Esc 返回，以及异步场景淡入淡出。")
    add_subheading(document, "正在处理")
    add_bullet(document, "第二、三关融合洞穴地图的开口、水面、水体范围、封边、碰撞体与出生点校准。")
    add_bullet(document, "教练与打捞目标模型的姿态、材质、比例、摆放和授权记录。")
    add_bullet(document, "水下音效、性能优化、完整中英文本与三关连续通关测试。")
    add_subheading(document, "主要阻碍与风险")
    add_bullet(document, "实景扫描模型存在封口面、破洞、薄壳与不规则坐标，容易导致水域不连通、穿模或出生在模型外。")
    add_bullet(document, "大体积 GLB、纹理和运行时导入会带来加载时间、内存与材质兼容风险；第三方模型授权需逐项核实。")

    add_heading(document, 3, "概念陈述")
    add_body(document, "这是一款单人第一人称洞穴潜水训练与打捞游戏。玩家从 60 m 训练池学习深度控制，跟随教练完成 Kelkaya 天坑下潜，并在一至两年后以成熟潜水员身份进入 Zacatón 洞穴执行遗体打捞；玩家依靠冷静操作、流程判断和对抗不确定水流完成任务，而不是依靠战斗。")

    add_heading(document, 4, "类型")
    add_bullet(document, "核心类型：第一人称 3D 潜水模拟 / 任务制冒险。")
    add_bullet(document, "辅助类型：训练教学、环境探索、轻度惊悚与叙事成长。")
    add_bullet(document, "不包含战斗、刷取或开放世界生存；MVP 强调短流程、清晰目标和真实水下控制感。")

    add_heading(document, 5, "目标受众")
    add_body(document, "目标玩家为 15 岁以上、对潜水、极限运动、实景探索、灾难救援或沉浸式模拟感兴趣的 PC 玩家。典型玩家可能接触过《Subnautica》《Narcosis》或写实探索类游戏，但不要求具备真实潜水知识。")
    add_body(document, "预期分级为 ESRB Teen（13+）：包含紧张的深水环境、缺氧式压迫氛围及非血腥遗体打捞主题，不表现肢解或写实血腥。")

    add_heading(document, 6, "概念段落与独特卖点")
    add_body(document, "游戏以“教练—学员—搭档”的关系串联三个阶段。玩家在第一关学习安全程序和深度控制；第二关由教练带领进入真实天坑，把训练迁移到开放环境；第三关时间跳跃一至两年，玩家已经技术成熟，与原教练并肩完成水下遗体打捞。每关都由简报、装备、入水、深度目标、返程和复盘构成，形成可理解且可重复的核心循环。")
    add_subheading(document, "独特卖点")
    add_bullet(document, "以真实洞穴/天坑扫描模型为基础，把实景地貌与可玩任务流程结合。")
    add_bullet(document, "从可控训练到真实救援的三段式成长叙事，教练既是教学者也是最终任务搭档。")
    add_bullet(document, "重力、浮力、水阻、深度压缩和会随机改变方向的水流共同制造持续但可读的水下不确定性。")
    add_bullet(document, "任务清单、深度读数、发光标记和颜色状态把复杂潜水流程转化为清晰反馈。")
    add_bullet(document, "不依赖战斗；压力来自深度、视野、路线、控制误差和打捞责任。")

    add_heading(document, 7, "玩家体验")
    add_body(document, "玩家扮演一名从初学者成长为洞穴潜水救援人员的潜水爱好者。MVP 为三关线性短篇，预计首次完整体验 30—45 分钟，单关约 8—15 分钟；解锁后可从任务地图重复挑战并刷新最佳时间。")
    add_bullet(document, "阶段一——学习：在安全环境中认识装备、教练、深度控制和水流干扰。")
    add_bullet(document, "阶段二——验证：进入 Kelkaya 天坑，在开放水域完成 10 m 稳定与 25 m 观察。")
    add_bullet(document, "阶段三——承担：在 Zacatón 深水洞穴搜索失踪潜水员，固定打捞带并拖带遗体返回。")
    add_body(document, "目标情绪曲线为：好奇与安全感 → 敬畏与紧张 → 责任、克制和任务完成后的释然。")

    add_heading(document, 8, "关键时刻")
    add_numbered_item(document, 1, "第一次入水", "完成教练说明并领取装备后，玩家主动跨过池边，移动模式从陆地切换为水下。")
    add_numbered_item(document, 2, "训练池下潜", "在随机水流干扰下分别于 20 m 停留 30 秒、40 m 停留 60 秒，并触碰 60 m 池底；完成目标后发光方块和圆环变灰停转。")
    add_numbered_item(document, 3, "真实天坑", "玩家在 Kelkaya 洞口看到实景地貌和单一连通水域，首次把训练中的深度控制用于开放环境。")
    add_numbered_item(document, 4, "一年后的重逢", "第三关开场通过简报说明玩家已经成长，原教练从指导者变为共同执行任务的搭档。")
    add_numbered_item(document, 5, "发现失踪者", "在深水路线尽头发现蜷缩的潜水员遗体，靠近并按 E 固定打捞带。")
    add_numbered_item(document, 6, "拖带返航", "携带目标上浮、离水并返回教练身边完成复盘，结算最终任务。")

    add_heading(document, 9, "美术、音效与音乐")
    add_subheading(document, "视觉方向")
    add_bullet(document, "环境：实景扫描与写实材质，洞外使用干燥岩地与植被，水下逐渐转为蓝灰、低对比、低能见度。")
    add_bullet(document, "角色：玩家不强调第三人称外观；教练使用正常站姿和专业户外/训练服，第三关失踪者采用自然蜷缩姿势。")
    add_bullet(document, "引导物：当前目标为青色发光，已完成目标变为灰色且停止旋转；危险深度使用克制的黄色提示。")
    add_bullet(document, "UI：半透明深色面板、白色标题、青色当前任务、绿色完成项、灰色未完成项。")
    add_subheading(document, "声音方向")
    add_bullet(document, "潜水过程以呼吸器、心跳、划水、气泡、水流、装备摩擦和洞穴混响为主。")
    add_bullet(document, "岸上使用环境风声与远处机械/自然底噪；菜单可使用低强度氛围音乐。")
    add_bullet(document, "正式下潜不铺设持续旋律，保留呼吸与水流作为节奏，关键目标仅使用短提示音。")

    add_heading(document, 10, "当前目标平台与系统要求")
    add_bullet(document, "平台：Windows PC；键盘 + 鼠标；单人离线。")
    add_bullet(document, "开发引擎：Unity 6.3 LTS（6000.3.10f1），Universal Render Pipeline。")
    add_bullet(document, "目标画面：1920 × 1080、60 FPS；最终最低/推荐硬件需在模型优化与完整构建后实测。")
    add_bullet(document, "暂不支持 VR、手柄、移动端与联网。")

    add_heading(document, 11, "竞品分析")
    add_numbered_item(document, 1, "Subnautica（2018）", "水下生存、探索与建造的代表作品。其核心是开放世界资源循环；本项目区别在于真实地点、短篇任务、教练程序和无制作系统。预算与收入不作为本项目已核实数据。")
    add_numbered_item(document, 2, "Narcosis（2017）", "第一人称深海叙事与生存恐怖。其压力主要来自叙事和恐怖事件；本项目更关注洞穴潜水训练、操作稳定性与非战斗救援。预算与收入资料未核实。")
    add_numbered_item(document, 3, "Cave Crave VR", "作为洞穴潜水/VR 体验方向的参考。其发行与财务信息待核实；本项目目前面向非 VR PC，以三关成长剧情、任务清单和实景天坑模型形成差异。")

    add_heading(document, 12, "盈利模式")
    add_body(document, "当前版本为课程/作品集用途的免费 Demo，不含广告、内购、抽卡或付费加速。若未来扩展为商业产品，建议采用一次性买断，后续地点作为内容扩展包；在玩法和性能验证前不进入商业化阶段。")

    add_heading(document, 13, "玩家目标与进程")
    add_subheading(document, "总体进程")
    add_body(document, "玩家按照 Level 1 → Level 2 → Level 3 的顺序解锁关卡。完成当前关卡后记录完成状态和最佳时间，并开放下一关；已解锁关卡可在三维任务地图中重玩。关卡内采用线性任务状态机，避免玩家跳过简报、装备或返航步骤。")
    add_subheading(document, "第一关：训练池")
    add_numbered_item(document, 1, "开始", "按 Enter 阅读训练目标并获得控制权。")
    add_numbered_item(document, 2, "教练与装备", "走到泳池另一侧按 E 对话，领取并检查潜水装备。")
    add_numbered_item(document, 3, "深度训练", "进入泳池；20 m 深度停留 30 秒；40 m 深度停留 60 秒；继续下潜并触碰 60 m 池底。")
    add_numbered_item(document, 4, "返程", "上浮、离水、登上池边，再次与教练对话并解锁第二关。")
    add_subheading(document, "第二关：Kelkaya 天坑")
    add_numbered_item(document, 1, "开始", "按 Enter，前往靠近洞口的教练位置完成安全简报。")
    add_numbered_item(document, 2, "准备", "在岸边领取潜水装备后进入洞内唯一连通水域。")
    add_numbered_item(document, 3, "开放水域训练", "下降至 10 m 完成稳定检查；下降至 25 m 并保持 15 秒完成观察。")
    add_numbered_item(document, 4, "返程", "上浮离水，返回教练身边完成复盘并解锁第三关。")
    add_subheading(document, "第三关：Zacatón 打捞")
    add_numbered_item(document, 1, "时间跳跃", "一至两年后，玩家与原教练共同执行遗体打捞。")
    add_numbered_item(document, 2, "准备", "复核打捞计划，领取带有固定装置的潜水装备并入水。")
    add_numbered_item(document, 3, "搜索", "沿洞穴路线下潜至约 260 m 搜索深度，接近 285 m 最后已知位置并寻找信标。")
    add_numbered_item(document, 4, "打捞", "发现蜷缩的失踪潜水员，按 E 固定打捞带，将目标拖带至水面。")
    add_numbered_item(document, 5, "结束", "把目标带回岸边教练位置，完成联合救援复盘和最终结算。")
    add_subheading(document, "核心循环与外层循环")
    add_body(document, "核心循环：阅读当前目标 → 观察环境与深度 → 控制方向/浮沉并抵消水流 → 到达目标区 → 保持或交互 → 获得颜色与勾选反馈 → 返回安全点。外层循环：完成任务 → 记录时间 → 解锁下一关/重玩 → 进入更深、更复杂且叙事责任更高的地点。")

    add_heading(document, 14, "游戏世界")
    add_body(document, "故事发生在现代，以专业洞穴潜水训练与救援为背景。玩家不是超能力角色，危险来自地形、水深、能见度与错误操作；教练的语言强调检查装备、保持冷静和及时修正。")
    add_bullet(document, "训练池：人工封闭环境，约 60 m 深，三个发光深度目标用于学习操作与水流补偿。")
    add_bullet(document, "Kelkaya 天坑：实景扫描的洞外与洞口环境，作为教练陪同的第一次真实开放水域训练。")
    add_bullet(document, "Zacatón 洞穴：与第二关共享融合地图骨架，但使用更深的原始洞穴路线和打捞布置。")
    add_bullet(document, "世界物理：岸上受重力、地面和墙体碰撞约束；水下受重力、浮力、水阻、深度密度、上浮阻力与随机水流共同影响。")
    add_bullet(document, "关卡推进为线性主线；已解锁地点可在任务地图自由重玩。MVP 不设置支线与彩蛋。")

    add_heading(document, 15, "用户界面")
    add_subheading(document, "控制方案")
    add_bullet(document, "W / A / S / D：前后左右移动；鼠标：自由观察，允许抬头看到正上方。")
    add_bullet(document, "Left Shift：加速；Space：岸上跳跃 / 水下上浮；C：水下下潜。")
    add_bullet(document, "E：与教练、装备、打捞目标交互；Enter：开始关卡。")
    add_bullet(document, "K：仅隐藏/恢复所有 HUD，不改变相机方向；M：打开三维任务地图；Esc：退出节点聚焦或返回原关卡。")
    add_subheading(document, "HUD 与流程")
    add_bullet(document, "左上任务面板显示关卡名和任务列表：当前项青色箭头、完成项绿色勾选、未完成项灰色圆点。")
    add_bullet(document, "水下阶段显示当前深度、目标深度、停留计时和 C/Space 操作提示。")
    add_bullet(document, "右上角显示小尺寸关卡计时；完成后显示通关时间、最佳时间、下一关、重玩和任务地图按钮。")
    add_bullet(document, "三维任务地图以美洲地形、发光节点和弯曲发光连线表示进程；支持鼠标点击、滚轮/按钮缩放和 Esc 返回。")
    add_subheading(document, "帮助、选项与保存")
    add_bullet(document, "帮助信息嵌入教练对话、任务面板和交互提示，不设置独立长教程页。")
    add_bullet(document, "进度、最高解锁关卡和最佳时间使用本地 PlayerPrefs 保存。MVP 暂无独立音量、画质和键位设置页。")

    add_heading(document, 16, "MVP 系统与功能")
    add_bullet(document, "移动系统：统一三关的陆地速度、冲刺、鼠标灵敏度、水下六向移动与碰撞停止。")
    add_bullet(document, "水下物理：浮力、重力、阻力、密度、表面阻尼、随机八方向水平水流及随机上下扰动。")
    add_bullet(document, "水域系统：每关定义水面高度、范围、底部与边界；玩家仅在有效水体内切换为水下控制。")
    add_bullet(document, "任务系统：每关独立状态机，严格控制简报、装备、目标、返航、复盘和完成顺序。")
    add_bullet(document, "交互与对话：距离判定 + E 键推进；对话期间锁定移动，结束后恢复。")
    add_bullet(document, "引导标记：发光方块/圆环旋转提示深度目标，完成后变灰停转；第三关使用搜索信标。")
    add_bullet(document, "装备门禁：没有领取潜水装备时禁止进入水域；装备点在岸边、靠近洞口但不与出生点重叠。")
    add_bullet(document, "关卡进度：顺序解锁、完成记录、最佳时间、任务地图状态与结算面板。")
    add_bullet(document, "场景系统：异步加载、淡入淡出、加载百分比、任务地图返回位置和场景重开。")
    add_bullet(document, "资源加载：使用 glTFast 导入 GLB/GLTF 实景模型与角色模型，并在运行时补充材质、碰撞与交互对象。")

    add_heading(document, 17, "游戏对象")
    add_bullet(document, "玩家：第一人称控制主体；状态包括岸上/水下、是否拥有装备、当前深度、速度和水流速度。")
    add_bullet(document, "教练：无攻击行为；提供开场简报、操作说明、安全提醒和结尾复盘。第一关是训练教练，第二关陪同开放水域，第三关成为救援搭档。")
    add_bullet(document, "潜水装备：岸边可交互拾取物，改变 PlayerGearState 并解除入水限制。第三关装备包含打捞固定装置。")
    add_bullet(document, "深度标记：20 m、40 m、60 m 训练目标及第二关 25 m 观察目标；具有发光、旋转、完成变灰状态。")
    add_bullet(document, "失踪潜水员：第三关打捞对象；初始为蜷缩姿势，玩家接近后按 E 固定，随后跟随玩家拖带至水面。")
    add_bullet(document, "水面/水体：负责视觉水面、水下状态确认、边界限制与深度计算。")
    add_bullet(document, "环境模型：训练池、Kelkaya 天坑、融合洞穴/Zacatón 深水路线；需要网格碰撞与防穿漏处理。")
    add_bullet(document, "系统对象：关卡任务控制器、计时与结算管理器、场景转换管理器、三维地图节点和 HUD。")

    add_heading(document, 18, "本地化")
    add_body(document, "首发目标为中文与英文。当前运行时任务和教练对话以英文硬编码为主，中文文案尚未接入统一本地化表，因此该项仍属于生产前工作。")
    add_bullet(document, "建议命名：`UI_`、`TASK_L1_`、`TASK_L2_`、`TASK_L3_`、`COACH_L1_` 等前缀 + 功能语义。")
    add_bullet(document, "变量使用占位符，例如 `{depth}`、`{seconds}`、`{best_time}`，避免在代码中拼接语言片段。")
    add_bullet(document, "最终字符串文件建议使用 UTF-8 CSV 或 Unity Localization String Table；数值与单位按语言格式化。")

    add_heading(document, 19, "工具")
    add_bullet(document, "Unity 6.3 LTS：场景、物理、UI、任务状态机、构建与性能分析。")
    add_bullet(document, "Blender：融合洞穴、删除封口面、打通水域、补洞、封边、法线、比例和导出 GLB。")
    add_bullet(document, "glTFast 6.19：运行时/编辑器导入 GLB 与 GLTF。")
    add_bullet(document, "URP 17.3、TextMeshPro/UGUI：材质、光照、任务 UI 和地图信息。")
    add_bullet(document, "Visual Studio / Rider 与 Git 或 Unity Version Control：脚本开发和版本管理。")
    add_bullet(document, "Sketchfab 等模型来源仅用于取得明确授权的资源，模型、作者、链接和许可证必须记录在 ATTRIBUTION/LICENCE 文件中。")

    add_heading(document, 20, "技术文档")
    add_subheading(document, "主要架构")
    add_bullet(document, "客户端单机架构，无服务器。Build Settings 包含 LevelSelect、Level1、Level2、Level3。")
    add_bullet(document, "SimpleMove/MouseLook 处理输入与移动；各 LevelMissionController 负责本关状态；LevelRunManager 负责计时结算；GameProgression 负责保存。")
    add_bullet(document, "SceneTransitionManager 常驻并异步切换场景；LevelSelectMap 在运行时生成任务地图；UnderwaterPanoramaMode 负责 K 键隐藏 HUD。")
    add_subheading(document, "项目约定")
    add_bullet(document, "场景文件：`Assets/Scenes/Level1.unity`、`Level2.unity`、`Level3.unity`、`LevelSelect.unity`。")
    add_bullet(document, "脚本按功能命名，关卡脚本使用 `LevelN` 前缀；共享大型模型放入 `Assets/StreamingAssets/Shared` 或 `Assets/Models/Shared`。")
    add_bullet(document, "模型优先使用 GLB；纹理采用 PNG/JPG；第三方资源同时保存授权与署名文本。")
    add_subheading(document, "自主与程序化系统")
    add_bullet(document, "随机水流按八个水平罗盘方向选择，并叠加随机上/下/中性分量；目标方向与强度按随机间隔改变。")
    add_bullet(document, "任务地图地形、节点、连线和 UI 在运行时生成。MVP 不含敌人 AI、寻路、经济或程序化关卡。")
    add_subheading(document, "主要技术风险与对策")
    add_bullet(document, "扫描网格开口被封：在 Blender 中删除遮挡面并建立一体化可潜水通道；在 Unity 中用可视水面和独立水体边界验证。")
    add_bullet(document, "碰撞过重或漏碰：使用简化碰撞网格/分区 MeshCollider，并在关键边缘增加 BoxCollider 防漏。")
    add_bullet(document, "大模型加载慢：合并材质、压缩纹理、拆分关卡模型并建立加载与内存基准。")
    add_bullet(document, "材质丢失：固定 URP 材质映射，保留 GLB 内嵌纹理并在导入后自动修复不支持的 shader。")
    add_bullet(document, "输入兼容：当前主要使用旧版 `Input` API，未来若迁移 Input System 必须一次性统一，避免双输入。")

    add_heading(document, 21, "想法与扩展")
    add_bullet(document, "加入气瓶压力、减压停留、照明电量、体温和更精细的浮力配重，但不影响当前 MVP。")
    add_bullet(document, "增加水下能见度变化、泥沙扰动、绳索引导、装备故障和动态水流事件。")
    add_bullet(document, "新增真实洞潜地点、分支救援任务、照片/勘探记录、教练评价等级与挑战模式。")
    add_bullet(document, "在稳定 PC 版本后评估手柄、VR、合作模式与无障碍选项。")

    add_heading(document, 22, "未决问题")
    add_subheading(document, "当前未解决")
    add_bullet(document, "第二、三关融合地图的洞口是否已真正打通，水面与可潜水体是否为同一连通体，仍需场景内完整下潜测试。")
    add_bullet(document, "第二、三关出生点、教练、装备和水边位置需在最终融合模型坐标确定后锁定。")
    add_bullet(document, "第三关失踪者姿态、拖带碰撞、视觉尺度和 285 m 深度节奏仍需可玩性验证。")
    add_bullet(document, "教练、潜水员、Kelkaya 与 Zacatón 模型的许可证、作者署名和是否允许改编需最终确认。")
    add_bullet(document, "声音资源、中文本地化、画质选项、最低硬件和最终构建大小尚未完成。")
    add_subheading(document, "已解决或已确定")
    add_bullet(document, "三关采用同一套移动逻辑、任务 UI、Enter 开始、K 隐藏 UI 和异步场景转换。")
    add_bullet(document, "第一关训练目标确定为 20 m / 30 秒、40 m / 60 秒、触碰 60 m 池底、上岸并二次对话。")
    add_bullet(document, "第二关是教练带队训练；第三关发生在一至两年后，主题为与原教练联合遗体打捞。")
    add_bullet(document, "第二、三关共享融合洞穴地图骨架，但关卡任务、深度、道具和叙事布置不同。")

    add_heading(document, 23, "原型")
    add_bullet(document, "Unity 工程：`D:\\Unity\\GP1`。")
    add_bullet(document, "主要场景：`Assets/Scenes/LevelSelect.unity`、`Level1.unity`、`Level2.unity`、`Level3.unity`。")
    add_bullet(document, "融合洞穴模型：`Assets/StreamingAssets/Shared/FusedCave/fused_cave_closed.glb`。")
    add_bullet(document, "可执行构建、演示视频、仓库地址、姓名、学号与课程信息：待填写。")

    document.save(working)

    # Repackage from the retained template and replace only document.xml. This
    # preserves the template's styles, numbering, settings, relationships and
    # other opaque parts byte-for-byte.
    with zipfile.ZipFile(source, "r") as source_zip, zipfile.ZipFile(working, "r") as work_zip:
        document_xml = work_zip.read("word/document.xml")
        with zipfile.ZipFile(package_tmp, "w") as output_zip:
            for info in source_zip.infolist():
                data = document_xml if info.filename == "word/document.xml" else source_zip.read(info.filename)
                output_zip.writestr(info, data)

    os.replace(package_tmp, output)
    working.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_document(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
